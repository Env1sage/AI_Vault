import io
import re
from collections.abc import Callable
from dataclasses import dataclass

import docx
import openpyxl
from pptx import Presentation
from pypdf import PdfReader

from vault_shared import ForbiddenError, get_logger
from vault_shared.connectors.google_drive import GoogleDriveClient
from vault_shared.db.models import ExtractionStatus, File

logger = get_logger("worker.enrichment.extraction")

# Skip extraction entirely above this size — bounds worker memory/time on a
# single file (Phase 5 spec's non-goal of full media processing, and the
# Handbook's general "never let one item's processing be unbounded").
_MAX_EXTRACTION_SIZE_BYTES = 20 * 1024 * 1024
# Extracted text is truncated to this many characters before storage — a
# pathologically text-dense file (e.g. a huge plain-text log) must not grow
# `file_extractions.extracted_text` without bound just because it slipped
# under the byte-size cap above.
_MAX_EXTRACTED_CHARS = 200_000
# openpyxl reads a spreadsheet lazily, but the text this produces is only
# meant to be a normalized "searchable attributes" surface (Phase 5 spec's
# Search Preparation), not a full data export — bounded to keep it small.
_MAX_XLSX_ROWS_PER_SHEET = 500

# Native Google Workspace types have no raw bytes — Drive must *export* them
# to a requested MIME type instead of a plain download.
_GOOGLE_NATIVE_EXPORTS: dict[str, tuple[str, str]] = {
    "application/vnd.google-apps.document": ("text/plain", "google_doc_export"),
    "application/vnd.google-apps.presentation": ("text/plain", "google_slides_export"),
    "application/vnd.google-apps.spreadsheet": ("text/csv", "google_sheets_export"),
}


def _extract_plain_text(content: bytes) -> str:
    return content.decode("utf-8", errors="replace")


def _extract_pdf(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(content: bytes) -> str:
    document = docx.Document(io.BytesIO(content))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _extract_xlsx(content: bytes) -> str:
    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
            if row_index >= _MAX_XLSX_ROWS_PER_SHEET:
                break
            lines.append(" ".join(str(cell) for cell in row if cell is not None))
    return "\n".join(lines)


def _extract_pptx(content: bytes) -> str:
    presentation = Presentation(io.BytesIO(content))
    lines: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.append(shape.text_frame.text)
    return "\n".join(lines)


# Binary formats downloaded via `GoogleDriveClient.download_file` (raw bytes).
_MIME_EXTRACTORS: dict[str, tuple[str, Callable[[bytes], str]]] = {
    "text/plain": ("plain_text", _extract_plain_text),
    "text/markdown": ("plain_text", _extract_plain_text),
    "application/pdf": ("pdf", _extract_pdf),
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": (
        "docx",
        _extract_docx,
    ),
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ("xlsx", _extract_xlsx),
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": (
        "pptx",
        _extract_pptx,
    ),
}

_CONTROL_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _sanitize(text: str) -> str:
    """Strips null/control bytes (Phase 5 spec's "sanitize extracted text
    before storage") and bounds length — never reject a file for this,
    just clean and truncate."""
    cleaned = _CONTROL_CHARS_RE.sub("", text)
    return cleaned[:_MAX_EXTRACTED_CHARS]


@dataclass(frozen=True)
class ExtractionOutcome:
    status: ExtractionStatus
    extractor_name: str | None
    text: str | None
    char_count: int | None
    error: str | None


class ContentExtractionService:
    """The Content Extraction Framework (Phase 5 spec) — the only place file
    *content* (not just metadata) is ever read, and only for formats with a
    registered extractor. Never writes back to Drive; never persists raw
    bytes anywhere, only the sanitized text a caller chooses to store.

    A network/auth failure talking to Drive (`DependencyUnavailableError`,
    `UnauthorizedError`) is allowed to propagate — if Drive is unreachable,
    every subsequent file will fail identically, so `EnrichmentService`
    treats it as a whole-job condition (same job-level-retry pattern as
    ADR-016's scanner). A parsing/format failure, or a `ForbiddenError`
    (a specific file's owner disabled download/copy/print for viewers —
    real, permanent, and unrelated to every other file) is caught here and
    reported as a normal outcome instead, so neither ever stops the rest of
    the enrichment run."""

    def __init__(self, drive_client: GoogleDriveClient) -> None:
        self._drive = drive_client

    def extract(self, *, access_token: str, file: File) -> ExtractionOutcome:
        if file.size_bytes is not None and file.size_bytes > _MAX_EXTRACTION_SIZE_BYTES:
            return ExtractionOutcome(
                status=ExtractionStatus.SKIPPED_TOO_LARGE,
                extractor_name=None,
                text=None,
                char_count=None,
                error=None,
            )

        mime_type = file.mime_type or ""

        if mime_type in _GOOGLE_NATIVE_EXPORTS:
            export_mime_type, extractor_name = _GOOGLE_NATIVE_EXPORTS[mime_type]
            try:
                content = self._drive.export_file(
                    access_token=access_token,
                    file_id=file.provider_file_id,
                    export_mime_type=export_mime_type,
                )
            except ForbiddenError:
                return self._forbidden(extractor_name)
            return self._finish(extractor_name, _extract_plain_text, content)

        if mime_type not in _MIME_EXTRACTORS:
            return ExtractionOutcome(
                status=ExtractionStatus.UNSUPPORTED,
                extractor_name=None,
                text=None,
                char_count=None,
                error=None,
            )

        extractor_name, extract_fn = _MIME_EXTRACTORS[mime_type]
        try:
            content = self._drive.download_file(
                access_token=access_token, file_id=file.provider_file_id
            )
        except ForbiddenError:
            return self._forbidden(extractor_name)
        return self._finish(extractor_name, extract_fn, content)

    @staticmethod
    def _forbidden(extractor_name: str) -> ExtractionOutcome:
        return ExtractionOutcome(
            status=ExtractionStatus.FORBIDDEN,
            extractor_name=extractor_name,
            text=None,
            char_count=None,
            error="The file owner has disabled downloading for this file.",
        )

    @staticmethod
    def _finish(
        extractor_name: str, extract_fn: Callable[[bytes], str], content: bytes
    ) -> ExtractionOutcome:
        try:
            text = _sanitize(extract_fn(content))
        except Exception as exc:  # noqa: BLE001 - any parser failure is a normal per-file outcome
            logger.warning(
                "content_extraction_failed", extra={"extractor": extractor_name, "error": str(exc)}
            )
            return ExtractionOutcome(
                status=ExtractionStatus.FAILED,
                extractor_name=extractor_name,
                text=None,
                char_count=None,
                error=str(exc)[:1024],
            )
        return ExtractionOutcome(
            status=ExtractionStatus.SUCCESS,
            extractor_name=extractor_name,
            text=text,
            char_count=len(text),
            error=None,
        )
