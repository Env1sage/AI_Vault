import io
import uuid

import docx
import openpyxl
from pptx import Presentation
from vault_shared import ForbiddenError
from vault_shared.db.models import ExtractionStatus, File
from worker.enrichment.extraction import ContentExtractionService


def _make_file(*, mime_type: str | None, size_bytes: int | None = 100) -> File:
    file = File(storage_source_id=uuid.uuid4(), provider_file_id="abc123")
    file.mime_type = mime_type
    file.size_bytes = size_bytes
    return file


class _FakeGoogleDriveClient:
    def __init__(
        self, *, content: bytes = b"", export_content: bytes = b"", forbidden: bool = False
    ) -> None:
        self.content = content
        self.export_content = export_content
        self.forbidden = forbidden
        self.download_calls = 0
        self.export_calls = 0

    def download_file(self, *, access_token: str, file_id: str) -> bytes:
        self.download_calls += 1
        if self.forbidden:
            raise ForbiddenError("Google Drive denied content access to this file.")
        return self.content

    def export_file(self, *, access_token: str, file_id: str, export_mime_type: str) -> bytes:
        self.export_calls += 1
        if self.forbidden:
            raise ForbiddenError("Google Drive denied content access to this file.")
        return self.export_content


def test_extract_skips_files_over_the_size_limit() -> None:
    service = ContentExtractionService(_FakeGoogleDriveClient())
    file = _make_file(mime_type="application/pdf", size_bytes=100 * 1024 * 1024)

    outcome = service.extract(access_token="token", file=file)

    assert outcome.status == ExtractionStatus.SKIPPED_TOO_LARGE
    assert outcome.text is None


def test_extract_reports_unsupported_for_an_unknown_mime_type() -> None:
    service = ContentExtractionService(_FakeGoogleDriveClient())
    file = _make_file(mime_type="application/octet-stream")

    outcome = service.extract(access_token="token", file=file)

    assert outcome.status == ExtractionStatus.UNSUPPORTED


def test_extract_plain_text_file() -> None:
    drive = _FakeGoogleDriveClient(content="hello world".encode())
    service = ContentExtractionService(drive)
    file = _make_file(mime_type="text/plain")

    outcome = service.extract(access_token="token", file=file)

    assert outcome.status == ExtractionStatus.SUCCESS
    assert outcome.text == "hello world"
    assert outcome.extractor_name == "plain_text"
    assert drive.download_calls == 1


def test_extract_docx_file() -> None:
    document = docx.Document()
    document.add_paragraph("Hello docx world")
    buffer = io.BytesIO()
    document.save(buffer)

    drive = _FakeGoogleDriveClient(content=buffer.getvalue())
    service = ContentExtractionService(drive)
    file = _make_file(
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    outcome = service.extract(access_token="token", file=file)

    assert outcome.status == ExtractionStatus.SUCCESS
    assert "Hello docx world" in outcome.text


def test_extract_xlsx_file() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["a", "b", "c"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    drive = _FakeGoogleDriveClient(content=buffer.getvalue())
    service = ContentExtractionService(drive)
    file = _make_file(
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    outcome = service.extract(access_token="token", file=file)

    assert outcome.status == ExtractionStatus.SUCCESS
    assert "a b c" in outcome.text


def test_extract_pptx_file() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "Hello pptx"
    buffer = io.BytesIO()
    presentation.save(buffer)

    drive = _FakeGoogleDriveClient(content=buffer.getvalue())
    service = ContentExtractionService(drive)
    file = _make_file(
        mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )

    outcome = service.extract(access_token="token", file=file)

    assert outcome.status == ExtractionStatus.SUCCESS
    assert "Hello pptx" in outcome.text


def test_extract_reports_failure_for_corrupt_content_without_raising() -> None:
    drive = _FakeGoogleDriveClient(content=b"not a real pdf")
    service = ContentExtractionService(drive)
    file = _make_file(mime_type="application/pdf")

    outcome = service.extract(access_token="token", file=file)

    assert outcome.status == ExtractionStatus.FAILED
    assert outcome.error is not None


def test_extract_google_native_document_uses_export() -> None:
    drive = _FakeGoogleDriveClient(export_content="Exported doc text".encode())
    service = ContentExtractionService(drive)
    file = _make_file(mime_type="application/vnd.google-apps.document")

    outcome = service.extract(access_token="token", file=file)

    assert outcome.status == ExtractionStatus.SUCCESS
    assert outcome.text == "Exported doc text"
    assert outcome.extractor_name == "google_doc_export"
    assert drive.export_calls == 1
    assert drive.download_calls == 0


def test_extract_reports_forbidden_when_the_owner_disabled_download() -> None:
    """Regression case: a shared file whose owner disabled download/copy/
    print for viewers returns 403 on content access even though its
    metadata was readable. This must be a normal per-file outcome, not a
    whole-job DependencyUnavailableError — retrying can never succeed."""
    drive = _FakeGoogleDriveClient(forbidden=True)
    service = ContentExtractionService(drive)
    file = _make_file(mime_type="application/pdf")

    outcome = service.extract(access_token="token", file=file)

    assert outcome.status == ExtractionStatus.FORBIDDEN
    assert outcome.error is not None


def test_extract_reports_forbidden_for_a_google_native_export_too() -> None:
    drive = _FakeGoogleDriveClient(forbidden=True)
    service = ContentExtractionService(drive)
    file = _make_file(mime_type="application/vnd.google-apps.document")

    outcome = service.extract(access_token="token", file=file)

    assert outcome.status == ExtractionStatus.FORBIDDEN
